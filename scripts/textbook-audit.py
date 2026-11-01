"""Build conservative, source-traceable textbook associations. No source translations are republished.
Usage: python3 scripts/textbook-audit.py /tmp/study-textbooks
Requires openpyxl for authoring only; the PWA consumes ordinary JSON.
"""
import sys,json,re,hashlib,unicodedata
from pathlib import Path
import openpyxl
root=Path(__file__).resolve().parents[1];source=Path(sys.argv[1])
items=sum([json.loads(p.read_text()) for p in (root/'quiz-data').glob('junior_*.json')],[])
lexical=[q for q in items if q.get('english')]
def en(s):return unicodedata.normalize('NFKC',str(s)).replace('’',"'").strip().casefold()
def ja(s):
 s=unicodedata.normalize('NFKC',s).strip()
 s=re.sub(r'^[~〜…]+(?:を|に|が|で|と|は)?','',s)
 return s.strip('。 .')
def senses(s):
 # Exact sense matches only; no synonym inference, no broad substring matching.
 return {ja(x) for x in re.split('[、・;；／/]',s) if ja(x)}
by={en(q['english']):q for q in lexical}
files={'nh.xlsx':'https://ten.tokyo-shoseki.co.jp/text/chu/eigo/download/documents/eigo_19goidata.xlsx','ow.xlsx':'https://www.kyoiku-shuppan.co.jp/textbook/chuu/eigo/files/r7_tangorengo-database_2603.xlsx','hwg.xlsx':'https://www.mitsumura-tosho.co.jp/download_file/view/9a89e7c9-55fb-42f6-9c98-e2c25522f08a/6046','mext.pdf':'https://www.mext.go.jp/content/20210531-mxt_kyoiku01-100002608_010.pdf','catalog.pdf':'https://www.mext.go.jp/content/20250415-mxt_kyokasyo02-000041714_2.pdf'}
for g in [1,2,3]:files[f'ow-plan{g}.xlsx']=f'https://www.kyoiku-shuppan.co.jp/textbook/chuu/eigo/files/r7eigo{g}_nenkeihyouka_2503.xlsx'
for g in [1,2,3]:files[f'blue-units{g}.pdf']=f'https://www.shinko-keirin.co.jp/keirinkan/chu_r7/english/file/english_tangenlist0{g}.pdf'
from datetime import date
checked=date.today().isoformat()
provenance=[dict(id=k,url=v,sha256=hashlib.sha256((source/k).read_bytes()).hexdigest(),retrieved=checked) for k,v in files.items()]
previousPath=root/'textbook-data/english-2026.json'
if previousPath.exists():
 previous=json.loads(previousPath.read_text());pinned={r['id']:r['sha256'] for r in previous['sources']}
 for row in provenance:
  if row['id'] in pinned and pinned[row['id']]!=row['sha256']:raise ValueError(f"Source changed: {row['id']}; review the edition, data revision and mapping before replacing its pinned checksum")
series=[('new-horizon','NEW HORIZON','東京書籍','002','wordlist_only'),('sunshine','Sunshine','開隆堂','009','source_unavailable'),('new-crown','NEW CROWN','三省堂','015','pending_source'),('one-world','ONE WORLD','教育出版','017','partial_words'),('here-we-go','Here We Go!','光村図書','038','partial_grammar'),('blue-sky','BLUE SKY','啓林館','061','curriculum_only')]
books=[]
for key,name,pub,code,status in series:
 for g in [1,2,3]:books.append(dict(id=f'{key}-{g}-2026',series=key,title=f'{name} {g}',publisher=pub,school='中学',subject='英語',year=g,usageYear=2026,editionYear=2025,bookNumber=f'{code}-{g+6}2',status=status,units=[]))
# Read only each file's current visible published planning sheet, never old/hidden drafts.
for g in [1,2,3]:
 sheet=openpyxl.load_workbook(source/f'ow-plan{g}.xlsx',data_only=True).worksheets[0]
 assert sheet.sheet_state=='visible' and sheet.title==f'年計案{g}年' and '令和7年度版' in str(sheet.cell(1,1).value)
 book=next(b for b in books if b['series']=='one-world' and b['year']==g);current=None
 for rowno,row in enumerate(sheet.values,1):
  if rowno<8:continue
  title,page=row[2],row[3]
  if title is not None:
   current=None
   if page is not None:
    first=str(title).split('\n')[0].strip();m=re.match(r'(Lesson|Reading|Further Reading|Springboard|Review Lesson|Project|Activities Plus|Useful Expressions|Tips for \w+|Sing Along|Classroom English)(?:\s*([0-9①②③④⑤]+))?',first)
    # Use short section identifiers, not textbook stories/activities or lesson prose.
    label=m.group(0).strip() if m else first.split('／')[0][:28]
    nums=re.findall(r'\d+',str(page))
    if nums:
     slug=re.sub(r'[^a-z0-9]+','-',unicodedata.normalize('NFKC',label).lower()).strip('-') or hashlib.sha256(label.encode()).hexdigest()[:8]
     current=dict(id=f'ow{g}-{slug}',label=label,pageStart=int(nums[0]),pageEnd=int(nums[-1]),source=f'ow-plan{g}.xlsx',row=rowno,items=[],listedEntries=0,unresolvedEntries=0)
     book['units'].append(current)
  if current and page is not None:
   nums=re.findall(r'\d+',str(page))
   if nums:current['pageStart']=min(current['pageStart'],int(nums[0]));current['pageEnd']=max(current['pageEnd'],int(nums[-1]))
 # Unit ranges must not overlap, and no page is assigned by guessing a neighboring unit.
 assert len({u['id'] for u in book['units']})==len(book['units'])
 for i,u in enumerate(book['units']):
  if i:assert book['units'][i-1]['pageEnd']<u['pageStart'],(g,u)
reviewPath=root/'textbook-data/one-world-reviewed.json'
review=json.loads(reviewPath.read_text())
assert review['version']==1 and review['source']=='ow.xlsx'
assert review['sourceSha256']==hashlib.sha256((source/'ow.xlsx').read_bytes()).hexdigest(), 'Stale review source'
reviews={r['row']:r for r in review['decisions']}
assert len(reviews)==len(review['decisions']), 'Duplicate review row'
consumed=set()
reviewMeta=dict(path=str(reviewPath.relative_to(root)),sha256=hashlib.sha256(reviewPath.read_bytes()).hexdigest(),approved=sum(r['decision']=='approved' for r in reviews.values()),held=sum(r['decision']=='hold' for r in reviews.values()))
wordreport=[];locations=[];seen_matches=set()
sheet=openpyxl.load_workbook(source/'ow.xlsx',data_only=True).worksheets[0]
assert '令和7年度版' in str(sheet.cell(2,1).value)
for rowno,row in enumerate(sheet.values,1):
 if rowno<8 or not isinstance(row[5],str):continue
 word=row[5];item=by.get(en(word));meaning=str(row[7] or '')
 match=bool(item and senses(item['japanese']) & senses(meaning))
 method='exact_normalized_sense'
 if rowno in reviews:
  decision=reviews[rowno];consumed.add(rowno)
  assert item and decision['sourceId']==item['id'] and decision['headword']==word, f'Stale review target: {rowno}'
  assert decision['publisherMeaningSha256']==hashlib.sha256(meaning.encode()).hexdigest(), f'Stale publisher sense: {rowno}'
  assert decision['materialSenseSha256']==hashlib.sha256((item['english']+'\0'+item['japanese']).encode()).hexdigest(), f'Stale material sense: {rowno}'
  assert decision['decision'] in ('approved','hold')
  match=decision['decision']=='approved'
  method='reviewed_equivalent_sense' if match else 'reviewed_hold'
 category='sense_confirmed' if match else 'sense_review' if item else 'headword_unmatched'
 form=re.search(r'([A-Za-z]+)\s*の(?:過去形|過去分詞|複数形|比較級|最上級)',meaning)
 reviewType='matched' if match else 'different_or_unconfirmed_sense' if item else 'inflection_reference' if form and en(form.group(1)) in by else 'multiword_or_expression' if re.search(r'\s',word.strip()) else 'proper_name_or_capitalized' if word[:1].isupper() else 'headword_candidate'

 locs=[]
 for cell in [row[2]]:
  if isinstance(cell,str):
   for grade,page in re.findall(r'([123])-(\d+)',cell):
    pair=(int(grade),int(page))
    if pair not in locs:locs.append(pair)
 mapped=0
 for g,page in locs:
  book=next(b for b in books if b['series']=='one-world' and b['year']==g)
  units=[u for u in book['units'] if u['pageStart']<=page<=u['pageEnd']]
  if not units:continue
  assert len(units)==1;unit=units[0];unit['listedEntries']+=1
  if not match:unit['unresolvedEntries']+=1;continue
  identity=(book['id'],unit['id'],item['id'])
  if identity not in seen_matches:
   unit['items'].append(dict(sourceId=item['id'],page=page,evidence=dict(source='ow.xlsx',sheet=sheet.title,row=rowno,method=method,meaningSha256=hashlib.sha256(meaning.encode()).hexdigest())))
   seen_matches.add(identity)
  mapped+=1
 wordreport.append(dict(headword=word,row=rowno,status=category,matchMethod=method if match or rowno in reviews else None,reviewType=reviewType,lemmaSourceId=by[en(form.group(1))]['id'] if form and en(form.group(1)) in by else None,sourceId=item['id'] if item else None,mappedLocations=mapped))
assert consumed==set(reviews), 'Unconsumed review rows'
# NH has no unit/page association and includes separate senses with the same spelling.
# Report headword overlap only: never claim meaning/unit confirmation from that file.
nh=[]
for sheet in openpyxl.load_workbook(source/'nh.xlsx',data_only=True).worksheets:
 if sheet.sheet_state!='visible':continue
 for rowno,row in enumerate(sheet.values,1):
  if rowno<2 or not isinstance(row[1],str):continue
  w=row[1];item=by.get(en(w));nh.append(dict(headword=w,sheet=sheet.title,row=rowno,status='headword_match_only' if item else 'headword_unmatched',sourceId=item['id'] if item else None))
# Grammar associations are explicitly reviewed, not inferred from keyword overlap.
grammarPath=root/'textbook-data/here-we-go-grammar-reviewed.json'
grammar=json.loads(grammarPath.read_text())
assert grammar['version']==1 and grammar['source']=='hwg.xlsx'
assert grammar['sourceSha256']==hashlib.sha256((source/'hwg.xlsx').read_bytes()).hexdigest(), 'Stale HWG source'
allBy={q['id']:q for q in items};unitKeys=set();hwg=openpyxl.load_workbook(source/'hwg.xlsx',data_only=True)
for reviewed in grammar['units']:
 g=reviewed['year'];number=reviewed['unitNumber'];key=(g,number)
 assert key not in unitKeys and g in (1,2,3) and number in range(1,9);unitKeys.add(key)
 sheet=hwg[reviewed['sheet']]
 assert sheet.sheet_state=='visible' and sheet.title==f'{g}年' and '令和7年度版' in ''.join(str(v or '') for v in next(sheet.values))
 row=list(sheet.values)[reviewed['row']-1]
 assert str(row[4]).splitlines()[0].strip()==f'Unit {number}', 'HWG unit moved'
 assert reviewed['languageMaterialSha256']==hashlib.sha256(str(row[8]).encode()).hexdigest(), 'HWG grammar changed'
 mapped=[];ids=set()
 for decision in reviewed['items']:
  item=allBy[decision['sourceId']]
  assert item['subject']=='英語' and item['field']=='英文法'
  assert item['id'] not in ids;ids.add(item['id'])
  assert decision['materialSha256']==hashlib.sha256(json.dumps(item,ensure_ascii=False,sort_keys=True,separators=(',',':')).encode()).hexdigest(), f"Stale HWG material: {item['id']}"
  mapped.append(dict(sourceId=item['id'],page=None,evidence=dict(source='hwg.xlsx',sheet=sheet.title,row=reviewed['row'],method='reviewed_grammar_topic',languageMaterialSha256=reviewed['languageMaterialSha256'])))
 book=next(b for b in books if b['series']=='here-we-go' and b['year']==g)
 book['units'].append(dict(id=f'hwg{g}-unit-{number}',label=f'Unit {number}',pageStart=None,pageEnd=None,source='hwg.xlsx',row=reviewed['row'],correspondenceType='grammar_topics',items=mapped,listedEntries=len(mapped),unresolvedEntries=0))
assert unitKeys=={(g,n) for g in (1,2,3) for n in range(1,9)}, 'Incomplete HWG unit review'
grammarMeta=dict(path=str(grammarPath.relative_to(root)),sha256=hashlib.sha256(grammarPath.read_bytes()).hexdigest())
hwgBooks=[b for b in books if b['series']=='here-we-go']
hwgReport=dict(scope='文法項目の単元対応のみ。語彙・読解・全用法の対応ではない。ページは未確認のため設定しない。',units=sum(len(b['units']) for b in hwgBooks),playableUnits=sum(bool(u['items']) for b in hwgBooks for u in b['units']),mappedMaterials=len({r['sourceId'] for b in hwgBooks for u in b['units'] for r in u['items']}),mappings=sum(len(u['items']) for b in hwgBooks for u in b['units']),reviewLedger=grammarMeta)
manifest=dict(version=1,usageYear=2026,editionYear=2025,checkedAt=checked,catalogSource='catalog.pdf',sources=provenance,reviewLedger=reviewMeta,grammarReviewLedger=grammarMeta,books=books,
 caveats=['出版社の監修・公認を意味しない。','語彙は確認できた語義・初出ページ、文法は確認できた単元項目に限定し、未対応内容を含む。','教科書本文・訳文・問題・音声は転載せず、問題文と解説は独自素材を使う。','初出箇所の対応であり、その単元の全既習語・全出現語一覧ではない。'])
(root/'textbook-data/english-2026.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2)+'\n')
report=dict(checkedAt=checked,scope='公開資料の掲載語／初出ページと既存素材の照合。未一致は未収録語確定ではなく、活用形・語義・固有名詞・表記等の要確認を含む。',sources=provenance,oneWorld=dict(totalEntries=len(wordreport),exactSenseConfirmed=sum(r['status']=='sense_confirmed' and r['matchMethod']=='exact_normalized_sense' for r in wordreport),reviewedSenseConfirmed=sum(r['matchMethod']=='reviewed_equivalent_sense' for r in wordreport),reviewedHeld=reviewMeta['held'],senseConfirmed=sum(r['status']=='sense_confirmed' for r in wordreport),senseReview=sum(r['status']=='sense_review' for r in wordreport),headwordUnmatched=sum(r['status']=='headword_unmatched' for r in wordreport),entries=wordreport),newHorizon=dict(totalEntries=len(nh),headwordMatched=sum(r['status']=='headword_match_only' for r in nh),entries=nh),hereWeGo=hwgReport)
(root/'docs/content/english-textbook-audit.json').write_text(json.dumps(report,ensure_ascii=False,indent=2)+'\n')
print('OW', {k:v for k,v in report['oneWorld'].items() if k!='entries'})
print('NH', {k:v for k,v in report['newHorizon'].items() if k!='entries'})
print('book units/matches',[(b['id'],len(b['units']),sum(len(u['items']) for u in b['units'])) for b in books if b['units']])
